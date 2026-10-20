"""Join official GSA metadata with the delivered cohort; never modify inputs.

Requires openpyxl. Outputs UTF-8 TSV and JSON, not rewritten workbooks.
"""
import argparse
import collections
import csv
import datetime as dt
import hashlib
import json
from pathlib import Path
import re
import shutil
import warnings

import openpyxl


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def workbook(path):
    # GSA exports advertise A1 dimensions even when the sheet has 477 rows.
    # Normal mode is intentional: trusting read_only dimensions truncates data.
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        book = openpyxl.load_workbook(path, read_only=False, data_only=True)
    sheets = {}
    for sheet in book:
        values = iter(sheet.values)
        header = [text(v) for v in next(values)]
        if len(set(header)) != len(header):
            raise ValueError(f"Duplicate headers in {sheet.title}")
        sheets[sheet.title] = [dict(zip(header, map(text, row))) for row in values
                               if any(text(v) for v in row)]
    book.close()
    return sheets


def csv_rows(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return [{k: text(v) for k, v in row.items()} for row in csv.DictReader(stream)]


def index(rows, key):
    result = {}
    for row in rows:
        value = row[key]
        if not value or value in result:
            raise ValueError(f"Missing or duplicate {key}: {value}")
        result[value] = row
    return result


def write_tsv(path, rows):
    if not rows:
        raise ValueError(f"Empty output: {path}")
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]), delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def fingerprint(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return {"path": str(path.resolve()), "bytes": path.stat().st_size, "sha256": digest.hexdigest()}


def compare_workbooks(current, previous):
    differences = []
    for name in sorted(set(current) | set(previous)):
        new = index(current.get(name, []), "Accession")
        old = index(previous.get(name, []), "Accession")
        for accession in sorted(set(new) | set(old)):
            a, b = new.get(accession, {}), old.get(accession, {})
            for field in sorted(set(a) | set(b)):
                if a.get(field) != b.get(field):
                    differences.append({"sheet": name, "accession": accession,
                                        "field": field, "previous": b.get(field, ""),
                                        "current": a.get(field, "")})
    return differences


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--official", type=Path, required=True)
    parser.add_argument("--cohort", type=Path, required=True)
    parser.add_argument("--previous", type=Path)
    parser.add_argument("--observations", type=Path, help="Previously recorded official page observations")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--project-root", type=Path, default=Path(__file__).resolve().parents[1])
    args = parser.parse_args()
    if args.output.exists():
        raise ValueError("Use a new snapshot directory; existing snapshots are preserved")
    official = workbook(args.official)
    samples = index(official["Sample"], "Accession")
    experiments = index(official["Experiment"], "Accession")
    runs = index(official["Run"], "Accession")
    cohort = index(csv_rows(args.cohort), "sample_id")
    conflicts, joined, files = [], [], []
    for run_id, run in sorted(runs.items()):
        experiment = experiments[run["Experiment accession"]]
        sample = samples[experiment["BioSample accession"]]
        for project in (run["BioProject accession"], experiment["BioProject accession"], sample["Project accession"]):
            if project != "PRJCA022804":
                raise ValueError(f"Unexpected project: {project}")
        match = re.fullmatch(r"Human gut metagenome from (AD|NC|MCI|SCD|SCS) sample", sample["Sample title"])
        if not match:
            raise ValueError(f"Unrecognized group title: {sample['Sample title']}")
        stage = match.group(1)
        local = cohort.get(run_id, {})
        row = {
            "study_accession": "CRA014435", "project_accession": "PRJCA022804",
            "run_accession": run_id, "experiment_accession": experiment["Accession"],
            "biosample_accession": sample["Accession"], "sample_alias": sample["Sample name"],
            "stage": stage, "stage_source": "NGDC Sample.Sample title",
            "sample_title_raw": sample["Sample title"], "host_disease_raw": sample["Host disease"],
            "host_subject_id_raw": sample["Host subject id"],
            "subject_id_status": "provided" if sample["Host subject id"] else "not_provided_in_ngdc_export",
            "host_age": sample["Host age"], "host_age_unit": sample["Host age unit"],
            "host_sex_raw": sample["Host sex"], "collection_date": sample["Collection date"],
            "geographic_location": sample["Geographic location"], "platform": experiment["Platform"],
            "library_strategy": experiment["Strategy"], "library_source": experiment["Source"],
            "library_layout": experiment["Layout"], "mate1_read_length_bp": experiment["Read length for mate1(bp)"],
            "mate2_read_length_bp": experiment["Read length for mate 2(bp)"],
            "included_in_current_delivery": "true" if local else "false",
            "scope_status": "in_current_delivery" if local else "not_in_current_delivery",
            "upstream_selection_reason": "pending_team_confirmation",
            "local_disease": local.get("disease", ""), "local_HPC_Batch": local.get("HPC_Batch", ""),
            "local_Group": local.get("Group", ""), "local_age": local.get("Age", ""),
            "local_gender": local.get("Gender", ""),
            "biosample_url": f"https://ngdc.cncb.ac.cn/biosample/browse/{sample['Accession']}",
        }
        if run["Run title"] != sample["Sample name"] or experiment["BioSample name"] != sample["Sample name"]:
            conflicts.append({"run_accession": run_id, "field": "official_alias_consistency"})
        if local:
            expected = {"Sample_name": sample["Sample name"], "Accession": sample["Accession"],
                        "disease": stage, "Age": sample["Host age"],
                        "Gender": {"male": "M", "female": "F"}.get(sample["Host sex"].lower(), sample["Host sex"])}
            for field, value in expected.items():
                if local.get(field) != value:
                    conflicts.append({"run_accession": run_id, "field": field,
                                      "local": local.get(field), "official": value})
        joined.append(row)
        for mate, name_column, md5_column, url_column in (
                (1, "Read filename 1", "Read file1 MD5", "DownLoad Read file1"),
                (2, "Read filename 2", "Read file2 MD5", "DownLoad  Read file2")):
            file_match = re.fullmatch(r"(.+) \((\d+) bytes\)", run[name_column])
            if not file_match or not re.fullmatch(r"[a-fA-F0-9]{32}", run[md5_column]):
                raise ValueError(f"Invalid file record: {run_id}, mate {mate}")
            filename, size = file_match.groups()
            url = run[url_column]
            if not url.startswith("ftp://download.big.ac.cn/gsa4/CRA014435/") or not url.endswith("/" + filename):
                raise ValueError(f"Unexpected official download URL: {url}")
            files.append({"run_accession": run_id, "biosample_accession": sample["Accession"],
                          "stage": stage, "mate": mate, "filename": filename, "size_bytes": int(size),
                          "md5_expected": run[md5_column], "official_download_url": url,
                          "included_in_current_delivery": row["included_in_current_delivery"],
                          "file_download_status": "not_downloaded", "md5_verification_status": "not_verified"})
    missing = sorted(set(cohort) - set(runs))
    selected = [r for r in joined if r["included_in_current_delivery"] == "true"]
    selected_files = [r for r in files if r["included_in_current_delivery"] == "true"]
    official_ad_nc = {r["run_accession"] for r in joined if r["stage"] in {"AD", "NC"}}
    matrix_checks = {}
    ko_path = args.project_root / "backend/storage/raw/ad-nc-ko-abundance/PRJCA022804/raw.csv"
    if ko_path.exists():
        ko = index(csv_rows(ko_path), "sample_id")
        storage_manifest = json.loads((args.project_root / "backend/storage_manifest.json").read_text(encoding="utf-8"))
        ko_config = next(d for d in storage_manifest["datasets"] if d["slug"] == "ad-nc-ko-abundance")
        label_map = ko_config["groupMapping"]
        matrix_checks["ko_group_mapping"] = label_map
        matrix_checks["ko_sample_set_matches"] = set(ko) == set(cohort)
        matrix_checks["ko_group_mismatches"] = sorted(k for k in set(ko) & set(cohort)
                                                     if label_map.get(ko[k]["label"], ko[k]["label"]) != cohort[k]["disease"])
    mag_path = args.project_root / "backend/storage/raw/source-metadata/PRJCA022804/metadata_NC_AD.csv"
    if mag_path.exists():
        mag = index(csv_rows(mag_path), "sample_id")
        matrix_checks["mag_metadata_exactly_matches"] = mag == cohort
    species_path = args.project_root / "backend/storage/raw/ad-nc-species/PRJCA022804/raw.xlsx"
    if species_path.exists():
        book = openpyxl.load_workbook(species_path, read_only=True, data_only=True)
        # Reading only the identity columns avoids loading thousands of abundance columns.
        values = book.active.iter_rows(min_row=2, max_col=3, values_only=True)
        species = [(text(r[0]), text(r[2])) for r in values if text(r[0])]
        book.close()
        species_ids = [sid for sid, _ in species]
        if len(set(species_ids)) != len(species_ids):
            raise ValueError("Duplicate sample IDs in species matrix")
        matrix_checks["species_total_samples"] = len(species)
        matrix_checks["species_prefix_counts"] = dict(collections.Counter(s[:3] for s in species_ids))
        matrix_checks["species_crr_set_matches"] = {s for s in species_ids if s.startswith("CRR")} == set(cohort)
        matrix_checks["species_group_mismatches"] = [s for s, group in species if s in cohort and group != cohort[s]["disease"]]
    differences = compare_workbooks(official, workbook(args.previous)) if args.previous else []
    comparison = {"changed_cells": len(differences),
                  "changed_fields": dict(collections.Counter(f"{r['sheet']}.{r['field']}" for r in differences))}
    checks = {
        "official_counts": {k: len(v) for k, v in official.items()},
        "official_stage_counts": dict(collections.Counter(r["stage"] for r in joined)),
        "cohort_count": len(cohort), "cohort_stage_counts": dict(collections.Counter(r["stage"] for r in selected)),
        "missing_current_runs": missing, "metadata_conflicts": conflicts,
        "current_equals_all_official_AD_NC_runs": set(cohort) == official_ad_nc,
        "unique_biosamples": len({r["biosample_accession"] for r in joined}),
        "unique_experiments": len({r["experiment_accession"] for r in joined}),
        "host_subject_id_empty_count": sum(not r["host_subject_id_raw"] for r in joined),
        "host_disease_empty_count": sum(not r["host_disease_raw"] for r in joined),
        "fastq_file_count": len(files), "fastq_current_file_count": len(selected_files),
        "fastq_total_bytes": sum(r["size_bytes"] for r in files),
        "fastq_current_bytes": sum(r["size_bytes"] for r in selected_files),
        "previous_export_comparison": comparison, "local_matrix_checks": matrix_checks,
    }
    validation_ok = (not missing and not conflicts and set(cohort) == official_ad_nc
                     and all(v for k, v in matrix_checks.items() if k.endswith("_matches"))
                     and not any(v for k, v in matrix_checks.items() if k.endswith("_mismatches")))
    checks["validation_ok"] = validation_ok
    args.output.mkdir(parents=True)
    raw = args.output / "raw"
    raw.mkdir()
    if args.observations:
        shutil.copy2(args.observations, args.output / "source_observations.json")
    source_records = []
    for label, path, target in [("official_ngdc", args.official, "CRA014435.xlsx"),
                                ("current_delivered_cohort", args.cohort, "metadata_NC_AD.csv"),
                                ("previous_user_export", args.previous, "CRA014435_previous.xlsx")]:
        if path:
            shutil.copy2(path, raw / target)
            source_records.append({"role": label, **fingerprint(path), "snapshot_path": f"raw/{target}"})
    write_tsv(args.output / "samples_all.tsv", joined)
    write_tsv(args.output / "samples_current185.tsv", selected)
    write_tsv(args.output / "fastq_files_all.tsv", files)
    write_tsv(args.output / "fastq_files_current185.tsv", selected_files)
    if differences:
        write_tsv(args.output / "official_export_changes.tsv", differences)
    write_json(args.output / "reconciliation.json", checks)
    manifest = {
        "schema_version": 1, "organized_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "study_accession": "CRA014435", "project_accession": "PRJCA022804",
        "official_export_source": "https://ngdc.cncb.ac.cn/gsa/browse/CRA014435",
        "official_export_download_method": "GSA page: metadata XLSX download control",
        "source_files": source_records,
        "transformation": {"join": "Run.Experiment accession -> Experiment.BioSample accession -> Sample.Accession",
                           "stage": "extracted from official Sample title; Host disease retained empty",
                           "sex_comparison": "male -> M; female -> F; raw values retained",
                           "current_scope": "membership in supplied metadata_NC_AD.csv; selection intent unconfirmed",
                           "no_subject_id_inference": True},
        "remaining_gaps": ["Host subject id is absent in the official export; request anonymized stable IDs or confirmation of one sample per participant from the team",
                           "Team must confirm why AD/NC was selected, exclusion/QC rules, and what HPC_Batch and Group mean",
                           "Actual input files, pipeline execution logs, tools/database versions, parameters and abundance normalization remain team deliverables"],
        "raw_fastq_downloaded": False,
        "organizer_script": fingerprint(Path(__file__)),
    }
    write_json(args.output / "provenance_manifest.json", manifest)
    total_tb = checks["fastq_total_bytes"] / 10**12
    current_tb = checks["fastq_current_bytes"] / 10**12
    group_lines = "\n".join(f"| {stage} | {count} |" for stage, count in checks["official_stage_counts"].items())
    report = f"""# CRA014435 元数据下载与核对

整理时间：{manifest['organized_at']}。本次只下载官方元数据，FASTQ 为下载清单，尚未下载或校验文件内容。

## 核对结果

- 官方 Sample / Experiment / Run 分别为 {len(samples)} / {len(experiments)} / {len(runs)} 条；逐级按 accession 关联。
- 当前交付清单 {len(cohort)} 条，其中 AD {checks['cohort_stage_counts'].get('AD', 0)}、NC {checks['cohort_stage_counts'].get('NC', 0)}。
- 当前清单与官方全部 AD/NC Run 集合一致：{checks['current_equals_all_official_AD_NC_runs']}。
- 本地样本别名、SAMC、分组、年龄、性别冲突：{len(conflicts)}；缺失的本地 Run：{len(missing)}。
- 完整校验通过：{validation_ok}。矩阵对照细节见 reconciliation.json。

| 官方标题中的分组 | Run 数 |
| --- | ---: |
{group_lines}

Host disease 和 Host subject id 各有 {checks['host_disease_empty_count']}、{checks['host_subject_id_empty_count']} 条空值。分组来自 Sample title，并单独记录出处；不改写原始空字段，也不把 Run、SAMC 或样本别名当成受试者 ID。官方一对一登记了样本、实验和 Run；受试者独立性仍需团队确认。

## 文件用途

| 文件 | 用途 |
| --- | --- |
| raw/CRA014435.xlsx | 本次 NGDC 原始下载，保留三个工作表 |
| raw/metadata_NC_AD.csv | 当前 185 个样本的团队交付清单快照 |
| raw/CRA014435_previous.xlsx | 用户此前提供的旧官方清单快照 |
| samples_all.tsv | 全部 476 个 Run 的编号、分组、元信息和当前纳入状态 |
| samples_current185.tsv | 当前 185 个 Run 的可关联元信息，前端接入优先使用 |
| fastq_files_all.tsv | 全队列 952 个原始文件的官方地址、字节大小和预期 MD5 |
| fastq_files_current185.tsv | 当前 185 个 Run 的 370 个双端文件清单 |
| official_export_changes.tsv | 新旧官方清单逐字段差异（有差异时生成） |
| reconciliation.json | 样本集合、冲突和本地矩阵对照结果 |
| provenance_manifest.json | 输入来源、SHA-256、整理规则及尚缺资料 |
| source_observations.json | 本次官方网页核查记录（另行保存） |

TSV 是 UTF-8，首行字段名，可用 Excel 打开。空值保持为空；Group、HPC_Batch 来自团队文件，含义尚待确认，不自行解释为临床分组或生物学批次。

KO 的分组核对采用 storage_manifest.json 中已有的 groupMapping（1=AD、0=NC）；不把数值标签直接与 AD/NC 文本比较。

## 原始测序文件与旧清单

全队列压缩 FASTQ 共 {len(files)} 个，合计 {checks['fastq_total_bytes']:,} 字节（约 {total_tb:.2f} TB，十进制）；当前子集 {len(selected_files)} 个，合计 {checks['fastq_current_bytes']:,} 字节（约 {current_tb:.2f} TB）。前端来源追溯使用上述元数据即可；生信组需要扩展分析时再按文件清单下载原始 reads。

新旧导出共 {comparison['changed_cells']} 个单元格发生变化，字段计数：{json.dumps(comparison['changed_fields'], ensure_ascii=False)}。下载地址以本次官方导出为准；是否能完整传输及 MD5 是否匹配须在实际 FASTQ 下载后检查。

## 还需要团队提供什么

1. 确认本次是否明确选择全部 AD/NC，为什么暂不纳入 MCI/SCD/SCS，以及是否另有质控或排除规则。
2. 受试者匿名稳定 ID；如果没有，请确认是否一人一份样本、有无重复采样。公开空字段无法靠编号推算补齐。
3. HPC_Batch、Group 的含义及产生方式；实际 shell 脚本、环境版本/运行日志、数据库版本、参数和丰度标准化说明。

这些文件已经补齐公开来源和编号关联；上游实际运行过程以及团队选择意图仍需团队说明。本次未改动分析矩阵、数据库或当前 185 个样本范围。

## 官方来源与论文

- [GSA CRA014435](https://ngdc.cncb.ac.cn/gsa/browse/CRA014435)：本次 XLSX 来源。
- [BioProject PRJCA022804](https://ngdc.cncb.ac.cn/bioproject/browse/PRJCA022804)：项目包含 CRA014435 和 476 个 BioSample。
- [2025 年来源论文](https://www.nature.com/articles/s41380-025-02973-7)：Metagenomic analysis characterizes stage-specific gut microbiota in Alzheimer’s disease；Data availability 明确指向 PRJCA022804。
- NGDC 当前还列有两篇 2026 年关联论文；这些关联不能替代团队实际数据生成方法。
"""
    (args.output / "README.md").write_text(report, encoding="utf-8")
    outputs = {p.name: fingerprint(p) for p in sorted(args.output.iterdir()) if p.is_file()}
    write_json(args.output / "output_checksums.json", outputs)
    print(json.dumps(checks, ensure_ascii=False, indent=2))
    if not validation_ok:
        raise SystemExit("Metadata reconciliation needs review; see reconciliation.json")


if __name__ == "__main__":
    main()
